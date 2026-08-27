from __future__ import annotations
from io import BytesIO
from openpyxl import load_workbook

from .intern_bezoek import _SECTIES


def _rows(ws, limit: int):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), 1):
        yield [str(c).strip().lower() if c is not None else "" for c in row]


def detect_kind(file_bytes: bytes, filename: str = "") -> str | None:
    """
    Herkent de twee bronbestanden die een eigen parser nodig hebben.

    Geeft 'videovisit', 'intern_bezoek' of None terug (None = gewoon dispatch-
    bestand). Leest in read_only-modus: het videobezoek-bestand is 12 MB met
    90 tabbladen, en dat openen op de gewone manier duurt zestien seconden.
    """
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return None

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Videobezoek: de kolomkop 'UUR VIDEO' is uniek genoeg.
            for low in _rows(ws, 8):
                if any("uur video" in c for c in low):
                    return "videovisit"

            # Intern bezoek: een cel die exact de sectiekop is. We eisen dat het
            # blad géén gewone dispatch-header heeft, zodat een dispatch-bestand
            # met 'intern bezoek' als bestemming hier niet per ongeluk in valt.
            heeft_sectie = False
            heeft_dispatch_header = False
            for low in _rows(ws, 30):
                if any(c in _SECTIES for c in low):
                    heeft_sectie = True
                if "naam" in low and ("bestemming" in low or "plaats" in low):
                    heeft_dispatch_header = True
            if heeft_sectie and not heeft_dispatch_header:
                return "intern_bezoek"
    finally:
        wb.close()

    return None
