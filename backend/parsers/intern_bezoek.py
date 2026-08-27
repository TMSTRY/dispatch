from __future__ import annotations
import re
from datetime import time
from io import BytesIO
from openpyxl import load_workbook
from .normalizer import normalize_cell
from .tijd import parse_time_loose

# Sectiekoppen in het bestand; de kop bepaalt meteen de bestemming.
_SECTIES = {
    "intern bezoek":     "Intern bezoek",
    "intern glasbezoek": "Intern glasbezoek",
}

_DAGEN = {
    "maandag": 0, "dinsdag": 1, "woensdag": 2, "donderdag": 3,
    "vrijdag": 4, "zaterdag": 5, "zondag": 6,
}

# Vangnet wanneer de regel bovenaan ontbreekt of anders geformuleerd wordt.
_STANDAARD_UUR   = time(16, 15)
_STANDAARD_DAGEN = [0, 2, 4]          # maandag, woensdag, vrijdag


def _clean_naam(value) -> str:
    """'CIL Aydin (ATV t.e.m. 26/08)' -> 'CIL Aydin'.

    De dienst noteert opmerkingen tussen haakjes achter de naam. Die horen niet
    in de zoeksleutel voor de celbezetting. Een ATV-vermelding is hier bewust
    geen reden om iemand over te slaan: net daarvoor dient het glasbezoek.
    """
    s = re.sub(r"\([^)]*\)", " ", str(value))
    return " ".join(s.split())


def _lees_instructieregel(ws) -> tuple[time, list[int]]:
    """Leest uur en dagen uit de regel bovenaan, bv.
    'Intern bezoek telkens op maandag - woensdag - vrijdag om 16h15'.
    Zo volgt de tool het bestand wanneer de dienst het uur of de dagen wijzigt."""
    uur = None
    dagen: list[int] = []

    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        for cell in row:
            if cell is None:
                continue
            tekst = str(cell)
            laag = tekst.lower()
            if laag.strip() in _SECTIES:
                continue                      # sectiekop, geen instructie
            for naam, nr in _DAGEN.items():
                if naam in laag and nr not in dagen:
                    dagen.append(nr)
            if uur is None and "om" in laag:
                uur = parse_time_loose(tekst)

    return (uur or _STANDAARD_UUR), (sorted(dagen) or _STANDAARD_DAGEN)


def parse_intern_bezoek(file_bytes: bytes, source_name: str = "intern bezoek") -> list[dict]:
    """
    Leest het intern-bezoekbestand.

    Elke datarij bevat TWEE gedetineerden die elkaar bezoeken: kolom A/B is de
    ene (mannelijke sectie), kolom C/D de andere (vrouwenafdeling). Allebei
    moeten ze op de dispatchlijst komen, dus elke rij levert twee regels op.

    Het bestand bevat geen datum — alleen de vaste dagen bovenaan. Elke rij
    krijgt daarom 'only_weekdays'; main.py laat ze bij het genereren weg wanneer
    de gekozen datum geen bezoekdag is.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    rows_out: list[dict] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        uur, dagen = _lees_instructieregel(ws)
        sectie: str | None = None

        for row in ws.iter_rows(values_only=True):
            if not row or all(v is None for v in row):
                continue

            # Sectiekop? Die staat samengevoegd over de volle breedte.
            eerste = str(row[0]).strip().lower() if row[0] is not None else ""
            if eerste in _SECTIES:
                sectie = _SECTIES[eerste]
                continue
            if sectie is None:
                continue                      # instructieregel of lege ruimte

            # Twee paren per rij: (cel, naam) links en (cel, naam) rechts.
            for cel_idx, naam_idx in ((0, 1), (2, 3)):
                naam_raw = row[naam_idx] if len(row) > naam_idx else None
                if naam_raw is None or not str(naam_raw).strip():
                    continue
                naam = _clean_naam(naam_raw)
                if not naam:
                    continue

                rows_out.append({
                    "uur":           uur,
                    "celnr":         normalize_cell(row[cel_idx] if len(row) > cel_idx else None),
                    "naam":          naam,
                    "voornaam":      None,
                    "bestemming":    sectie,
                    "source":        source_name,
                    "only_weekdays": dagen,
                })

    return rows_out
