from __future__ import annotations
from datetime import datetime
from io import BytesIO
from openpyxl import load_workbook
from .normalizer import normalize_cell
from .tijd import parse_time_loose

# Terugkerende afspraken per weekdag met andere inrichtingen. Op vraag van de
# dienst blijft dit tabblad buiten de dispatchlijst: de tweewekelijkse ritmes
# zijn niet af te leiden zonder ijkpunt en de helft mist een celnummer.
_SKIP_SHEETS = {"bez ander inrichting"}

BESTEMMING = "Videobezoek"


def _find_header(ws):
    """
    Zoekt de headerrij en leidt de kolomposities af uit de koppen zelf.

    De koppen verschillen per tabblad — 'Naam gedetineerde', 'Naam gedetinneerde'
    (dubbele n), 'Gedetineerden', en 'DATUM ' / 'Datum Video' / 'datum'. Vaste
    kolomnummers zouden dus op een deel van de bladen stilzwijgend fout zitten.
    """
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
        low = [str(c).strip().lower() if c is not None else "" for c in row]
        if not any("uur video" in c for c in low):
            continue
        pos: dict[str, int] = {}
        for j, c in enumerate(low):
            if "celnr" in c or c == "cel":
                pos.setdefault("cel", j)
            elif "gedetin" in c:
                pos.setdefault("naam", j)
            elif c.startswith("datum"):
                pos.setdefault("datum", j)
            elif "uur video" in c:
                pos.setdefault("uur", j)
        return i, pos
    return None, None


def _cell(row, idx):
    return row[idx] if idx is not None and len(row) > idx else None


def parse_videovisit(file_bytes: bytes, source_name: str = "videobezoek") -> list[dict]:
    """
    Leest het videobezoek-bestand (VideoVisit) van de bezoekdienst.

    Het bestand is een lopend archief met per week een tabblad. Van de ~148.000
    gevulde rijen is maar een fractie een echte afspraak: alleen rijen met een
    ingevulde datum tellen. De rest is een register van goedgekeurde bezoekers
    — wie bij wie op bezoek mag — en hoort niet op de dispatchlijst.

    Elke rij krijgt een 'visit_date'; main.py houdt bij het genereren alleen de
    afspraken over die op de gekozen datum vallen. Datumfiltering gebeurt over
    alle tabbladen heen in plaats van via de tabbladnaam: die namen bevatten
    geen jaartal, en enkele datums blijken in het verkeerde weekblad te staan.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    rows_out: list[dict] = []
    seen: set[tuple] = set()

    try:
        for sheet_name in wb.sheetnames:
            if sheet_name.strip().lower() in _SKIP_SHEETS:
                continue

            ws = wb[sheet_name]
            header_row, pos = _find_header(ws)
            # Zonder datum- of uurkolom valt er niets te plannen.
            if header_row is None or "datum" not in pos or "uur" not in pos:
                continue

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if not row:
                    continue

                datum = _cell(row, pos["datum"])
                if not isinstance(datum, datetime):
                    continue          # registerrij zonder afspraak

                naam_raw = _cell(row, pos.get("naam"))
                if naam_raw is None or not str(naam_raw).strip():
                    continue
                naam = " ".join(str(naam_raw).split())

                uur = parse_time_loose(_cell(row, pos["uur"]))
                cel = normalize_cell(_cell(row, pos.get("cel")))

                # Dezelfde afspraak staat soms in twee weekbladen. Zelfde persoon
                # op zelfde dag en uur is een dubbel; een tweede gesprek later op
                # de dag heeft een ander uur en blijft dus wel staan.
                key = (naam.lower(), datum.date(), uur)
                if key in seen:
                    continue
                seen.add(key)

                rows_out.append({
                    "uur":        uur,
                    "celnr":      cel,
                    # Voluit doorgeven: de matcher plakt naam en voornaam toch weer
                    # aan elkaar voor de zoeksleutel, en haalt de nette schrijfwijze
                    # daarna uit de celbezetting.
                    "naam":       naam,
                    "voornaam":   None,
                    "bestemming": BESTEMMING,
                    "source":     source_name,
                    "visit_date": datum.date(),
                })
    finally:
        wb.close()

    return rows_out
